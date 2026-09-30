"""Focused Step 1 acceptance tests for generic ingestion and provenance."""

import io
import os
import sys
from unittest.mock import patch

import fitz
import httpx
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from database import Base
from app.models.official_source import OfficialSource, OfficialDocument
from app.models.document import Document
from app.services.document_models import DocumentResult
from app.services.ingestion_service import calculate_sha256
from app.services.official_source_connector import MinistryOfCoalConnector, OfficialSourceConnector
from app.services.official_sync_service import sync_official_source
from app.services.parsing_service import classify_pdf_page_signals, parse_csv_result, parse_docx_result, parse_pdf_result, parse_xlsx_result
from app.services.processing_states import can_transition, transition


def test_page_classifier_is_independent_and_selective():
    assert classify_pdf_page_signals(800, 4, 0, 0.0)[0:2] == ("DIGITAL", False)
    assert classify_pdf_page_signals(0, 0, 1, 0.95)[0:2] == ("SCANNED", True)
    assert classify_pdf_page_signals(40, 2, 1, 0.55)[0:2] == ("MIXED", True)
    assert classify_pdf_page_signals(250, 3, 1, 0.55)[0:2] == ("MIXED", True)


def test_csv_preserves_tabular_cell_provenance():
    result = parse_csv_result("report.csv", b"Subsidiary,April,May\nECL,7.07,8.10\n")
    table = result.tables[0]
    assert table.headers == ["Subsidiary", "April", "May"]
    assert table.rows == [["ECL", "7.07", "8.10"]]
    assert table.cells[3]["coordinate"] == "A2"


def test_xlsx_preserves_sheets_merges_formulas_and_values(tmp_path):
    import openpyxl
    workbook = openpyxl.Workbook()
    production = workbook.active
    production.title = "Production"
    production.append(["Subsidiary", "April", "Total"])
    production.append(["ECL", 7.07, "=B2"])
    production.merge_cells("A4:B4")
    dispatch = workbook.create_sheet("Despatch")
    dispatch.append(["Subsidiary", "April"])
    dispatch.append(["ECL", 6.5])
    output = io.BytesIO()
    workbook.save(output)
    result = parse_xlsx_result("stats.xlsx", output.getvalue())
    assert result.metadata["sheet_names"] == ["Production", "Despatch"]
    assert result.tables[0].sheet_name == "Production"
    assert "A4:B4" in result.tables[0].merged_cells
    assert result.tables[0].formulas["C2"] == "=B2"
    assert result.tables[1].sheet_name == "Despatch"


def test_docx_preserves_paragraph_table_paragraph_order():
    from docx import Document
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "ECL"
    table.cell(1, 1).text = "7.07"
    document.add_paragraph("After table")
    output = io.BytesIO()
    document.save(output)
    result = parse_docx_result("report.docx", output.getvalue())
    page_text = result.pages[0].text
    assert page_text.index("Before table") < page_text.index("Header") < page_text.index("After table")


def test_mixed_pdf_ocr_only_runs_for_required_page():
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((40, 60), "This is a fully digital native page with enough meaningful text to classify it reliably.")
    image = Image.new("RGB", (500, 500), "white")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    scanned = pdf.new_page()
    scanned.insert_image(scanned.rect, stream=image_bytes.getvalue())
    raw = pdf.tobytes()
    calls = []

    def fake_ocr(image):
        calls.append(1)
        return "Scanned evidence 781.05", [], 0.96

    with patch("app.services.parsing_service._ocr_image", side_effect=fake_ocr):
        result = parse_pdf_result("mixed.pdf", raw)
    assert [page.classification for page in result.pages] == ["DIGITAL", "SCANNED"]
    assert calls == [1]
    assert result.pages[1].text == "Scanned evidence 781.05"


def test_official_connector_discovers_bounded_same_host_resources():
    html = '<a href="/files/stats.pdf">Monthly Statistics</a><a href="https://example.com/other.pdf">Ignore</a><a href="/about">About</a>'
    response = httpx.Response(200, text=html, request=httpx.Request("GET", "https://coal.nic.in/major-statistics-page"))

    class FakeClient:
        def get(self, url):
            return response

    docs = MinistryOfCoalConnector(client=FakeClient()).discover_documents()
    assert len(docs) == 1
    assert docs[0].url == "https://coal.nic.in/files/stats.pdf"
    assert docs[0].title == "Monthly Statistics"


def test_state_machine_rejects_invalid_transition():
    class Doc:
        processing_status = "DISCOVERED"
        processing_warnings = []

    doc = Doc()
    assert can_transition("DISCOVERED", "DOWNLOADED")
    assert not can_transition("DISCOVERED", "READY")
    transition(doc, "DOWNLOADED")
    try:
        transition(doc, "READY")
    except ValueError:
        pass
    else:
        raise AssertionError("invalid state transition was accepted")


def test_official_sync_is_idempotent_and_versions_changed_content():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    source = OfficialSource(name="Test Ministry", organization="Test Ministry", base_url="https://coal.nic.in/major-statistics-page")
    session.add(source)
    session.commit()
    state = {"content": b"Subsidiary,April\nECL,7.07\n"}
    page = '<a href="/files/stats.csv">Stats</a>'

    class FakeResponse:
        def __init__(self, content, text=None):
            self.content = content
            self.text = text if text is not None else content.decode("utf-8")
            self.headers = {"content-length": str(len(content))}
        def raise_for_status(self):
            return None

    class FakeClient:
        def get(self, url):
            return FakeResponse(page.encode() if url.endswith("major-statistics-page") else state["content"], page if url.endswith("major-statistics-page") else None)

    connector = OfficialSourceConnector("https://coal.nic.in/major-statistics-page", "Test Ministry", client=FakeClient())
    first = sync_official_source(session, source, connector)
    second = sync_official_source(session, source, connector)
    state["content"] = b"Subsidiary,April\nECL,8.10\n"
    third = sync_official_source(session, source, connector)
    assert first["new"] == 1
    assert second["unchanged"] == 1
    assert third["updated"] == 1
    versions = session.query(OfficialDocument).filter(OfficialDocument.source_id == source.id).order_by(OfficialDocument.version).all()
    assert [item.version for item in versions] == [1, 2]
    assert session.query(Document).count() == 2


def test_official_sync_does_not_report_download_only_document_as_ingested():
    """Regression: official sync must invoke the common processing pipeline."""
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    source = OfficialSource(name="Test Ministry", organization="Test Ministry", base_url="https://coal.nic.in/major-statistics-page")
    session.add(source)
    session.commit()

    class Item:
        url = "https://coal.nic.in/files/live.pdf"
        title = "live.pdf"
        category = "Statistics"
        publication_date = None

    class Connector:
        def discover_documents(self):
            return [Item()]

        def download_document(self, item):
            return b"%PDF-1.7 live document bytes"

        def extract_source_metadata(self, item):
            return {"source_organization": "Test Ministry", "title": item.title}

    calls = []

    def fake_process(db, file_bytes, original_filename, user_id, **kwargs):
        document = Document(
            filename=original_filename,
            file_path="fixture/live.pdf",
            file_hash=calculate_sha256(file_bytes),
            file_type="PDF",
            file_size_bytes=len(file_bytes),
            status="PENDING",
            processing_status="DOWNLOADED",
            source_type="OFFICIAL",
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        return document

    def fake_pipeline(db, document_id):
        calls.append(document_id)
        document = db.get(Document, document_id)
        document.status = "PARSED"
        document.processing_status = "READY"
        db.commit()
        return True

    with patch("app.services.official_sync_service.process_file_ingestion", side_effect=fake_process), patch(
        "app.services.official_sync_service.execute_document_processing_pipeline", side_effect=fake_pipeline
    ):
        result = sync_official_source(session, source, Connector())

    document = session.query(Document).one()
    record = session.query(OfficialDocument).one()
    assert calls == [document.id]
    assert result["new"] == 1
    assert result["failed"] == 0
    assert record.download_status == "INGESTED"
    assert document.processing_status == "READY"


def test_official_sync_marks_pipeline_failure_without_false_ingested_state():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    source = OfficialSource(name="Test Ministry", organization="Test Ministry", base_url="https://coal.nic.in/major-statistics-page")
    session.add(source)
    session.commit()

    class Item:
        url = "https://coal.nic.in/files/broken-live.pdf"
        title = "broken-live.pdf"
        category = "Statistics"
        publication_date = None

    class Connector:
        def discover_documents(self):
            return [Item()]

        def download_document(self, item):
            return b"downloaded bytes"

        def extract_source_metadata(self, item):
            return {"source_organization": "Test Ministry", "title": item.title}

    def fake_process(db, file_bytes, original_filename, user_id, **kwargs):
        document = Document(
            filename=original_filename,
            file_path="fixture/broken-live.pdf",
            file_hash=calculate_sha256(file_bytes),
            file_type="PDF",
            file_size_bytes=len(file_bytes),
            status="PENDING",
            processing_status="DOWNLOADED",
            source_type="OFFICIAL",
        )
        db.add(document)
        db.commit()
        db.refresh(document)
        return document

    with patch("app.services.official_sync_service.process_file_ingestion", side_effect=fake_process), patch(
        "app.services.official_sync_service.execute_document_processing_pipeline", return_value=False
    ):
        result = sync_official_source(session, source, Connector())

    record = session.query(OfficialDocument).one()
    assert result["failed"] == 1
    assert record.download_status == "PROCESSING_FAILED"
    assert record.document_id is not None


def test_official_connector_crawls_category_pages_and_records_partial_page_failure():
    pages = {
        "https://coal.nic.in/major-statistics-page": '<a href="/major-statistics/production">Production</a><a href="/major-statistics/broken">Broken</a><a href="https://evil.example/file.pdf">Ignore</a>',
        "https://coal.nic.in/major-statistics/production": '<a href="/files/production.pdf">Production Report</a>',
    }

    class Response:
        status_code = 200
        headers = {}
        url = None
        def __init__(self, text):
            self.text = text
        def raise_for_status(self):
            return None

    class Client:
        def get(self, url, **kwargs):
            if url.endswith("/broken"):
                raise httpx.ConnectTimeout("category unavailable")
            return Response(pages[url])

    connector = OfficialSourceConnector(
        "https://coal.nic.in/major-statistics-page", "Ministry of Coal", client=Client(),
        max_attempts=2, backoff_factor=0, sleep=lambda _: None,
    )
    report = connector.discover_documents_with_report()
    assert [item.url for item in report.documents] == ["https://coal.nic.in/files/production.pdf"]
    assert report.pages_checked == 2
    assert len(report.page_failures) == 1
    assert report.page_failures[0]["url"].endswith("/broken")


def test_official_connector_retries_transient_status_and_keeps_host_restriction():
    responses = [
        httpx.Response(503, request=httpx.Request("GET", "https://coal.nic.in/major-statistics-page")),
        httpx.Response(200, text='<a href="/files/stats.pdf">Stats</a>', request=httpx.Request("GET", "https://coal.nic.in/major-statistics-page")),
    ]
    calls = []

    class Client:
        def get(self, url, **kwargs):
            calls.append(url)
            return responses.pop(0)

    connector = OfficialSourceConnector(
        "https://coal.nic.in/major-statistics-page", "Ministry of Coal", client=Client(),
        max_attempts=2, backoff_factor=0, sleep=lambda _: None,
    )
    assert len(connector.discover_documents()) == 1
    assert len(calls) == 2
    assert connector._is_allowed_url("https://coal.nic.in/files/stats.pdf")
    assert not connector._is_allowed_url("https://example.com/files/stats.pdf")


def test_sync_partial_document_failure_preserves_success_and_registry_audit():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    session = sessionmaker(bind=engine)()
    source = OfficialSource(name="Test Ministry", organization="Test Ministry", base_url="https://coal.nic.in/major-statistics-page")
    session.add(source)
    session.commit()

    class Item:
        def __init__(self, url, title):
            self.url, self.title, self.category, self.publication_date = url, title, "Statistics", None

    good = Item("https://coal.nic.in/files/good.csv", "Good.csv")
    broken = Item("https://coal.nic.in/files/broken.csv", "Broken.csv")

    class Connector:
        def discover_documents(self):
            return [good, broken]
        def download_document(self, item):
            if item is broken:
                raise httpx.ConnectTimeout("document unavailable")
            return b"Subsidiary,April\nECL,7.07\n"
        def extract_source_metadata(self, item):
            return {"source_organization": "Test Ministry", "title": item.title}

    result = sync_official_source(session, source, Connector())
    assert result["new"] == 1
    assert result["failed"] == 1
    assert result["status"] == "PARTIAL"
    records = session.query(OfficialDocument).order_by(OfficialDocument.document_url).all()
    assert [record.download_status for record in records] == ["FAILED", "INGESTED"]
    assert session.query(Document).count() == 1
