"""Task 3 acceptance tests for generic table structure and provenance.

These tests intentionally assert source structure, not coal-domain meaning.
"""

import csv
import io
import os
import sys

import fitz
import openpyxl
import pytest
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.services.parsing_service import parse_document_result, parse_pdf_result


def _bordered_pdf(pages):
    document = fitz.open()
    for page_spec in pages:
        page = document.new_page(width=500, height=360)
        x_edges = page_spec.get("x_edges", [40 + index * 140 for index in range(len(page_spec["rows"][0]) + 1)])
        y_edges = page_spec.get("y_edges", [40 + index * 35 for index in range(len(page_spec["rows"]) + 1)])
        for x in x_edges:
            page.draw_line((x, y_edges[0]), (x, y_edges[-1]))
        for y in y_edges:
            page.draw_line((x_edges[0], y), (x_edges[-1], y))
        for row_index, row in enumerate(page_spec["rows"]):
            for column_index, value in enumerate(row):
                if value != "":
                    page.insert_text(
                        (x_edges[column_index] + 3, y_edges[row_index + 1] - 9),
                        str(value),
                        fontsize=9,
                    )
        for text, y in page_spec.get("extra_text", []):
            page.insert_text((45, y), text, fontsize=8)
    return document.tobytes()


def _borderless_pdf():
    document = fitz.open()
    page = document.new_page(width=500, height=360)
    rows = [["Metric", "April", "May"], ["Production", "781.05", "781.50"], ["Change (%)", "0.781", "78.105"]]
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            page.insert_text((50 + column_index * 150, 60 + row_index * 35), value, fontsize=10)
    return document.tobytes()


def test_native_pdf_cells_keep_exact_values_and_geometry():
    pdf = _bordered_pdf([{"rows": [["Metric (MT)", "April", "May"], ["Production", "781.05", "781.50"], ["Blank", "", "-5"]]}])
    result = parse_pdf_result("", pdf, filename="numeric.pdf")
    assert len(result.tables) == 1
    table = result.tables[0]
    assert table.headers == ["Metric (MT)", "April", "May"]
    assert table.rows[1] == ["Blank", "", "-5"]
    assert len(table.cells) == 9
    value_cell = next(cell for cell in table.cells if cell["value"] == "781.05")
    assert value_cell["row_index"] == 1
    assert value_cell["column_index"] == 1
    assert value_cell["bounding_box"] == [180.0, 75.0, 320.0, 110.0]
    assert next(cell for cell in table.cells if cell["value"] == "")["bounding_box"] is not None
    assert table.extraction_method == "NATIVE_TABLE"


def test_borderless_native_table_uses_bounded_text_strategy():
    result = parse_pdf_result("", _borderless_pdf(), filename="borderless.pdf")
    assert len(result.tables) == 1
    table = result.tables[0]
    assert table.headers == ["Metric", "April", "May"]
    assert ["Production", "781.05", "781.50"] in table.rows
    assert "text strategy" in " ".join(table.warnings)
    assert all(cell["extraction_method"] == "NATIVE_TABLE" for cell in table.cells)


def test_multi_page_tables_keep_page_and_repeated_header_provenance():
    rows = [["Metric", "Value"], ["Production", "781.05"]]
    result = parse_pdf_result("", _bordered_pdf([{"rows": rows}, {"rows": rows}]), filename="multi.pdf")
    assert [(table.page_number, table.table_number) for table in result.tables] == [(1, 1), (2, 1)]
    assert all(table.headers == ["Metric", "Value"] for table in result.tables)
    assert result.tables[0].cells[-1]["bounding_box"] != result.tables[1].cells[-1]["bounding_box"] or result.tables[0].page_number != result.tables[1].page_number


def test_multiple_tables_on_page_are_separate_generic_tables():
    document = fitz.open()
    page = document.new_page(width=600, height=400)
    for x0, y0, title, value in [(40, 40, "A", "781.05"), (330, 200, "B", "0.781")]:
        page.draw_rect((x0, y0, x0 + 200, y0 + 70))
        page.draw_line((x0 + 100, y0), (x0 + 100, y0 + 70))
        page.draw_line((x0, y0 + 35), (x0 + 200, y0 + 35))
        page.insert_text((x0 + 4, y0 + 25), title)
        page.insert_text((x0 + 104, y0 + 60), value)
    result = parse_pdf_result("", document.tobytes(), filename="two-tables.pdf")
    assert len(result.tables) == 2
    assert [table.table_number for table in result.tables] == [1, 2]
    assert {table.cells[-1]["value"] for table in result.tables} == {"781.05", "0.781"}


def test_footnotes_and_units_remain_text_not_domain_interpretation():
    pdf = _bordered_pdf([{
        "rows": [["Production (MT)", "April"], ["Total", "781.05"]],
        "x_edges": [40, 260, 460],
        "y_edges": [80, 115, 150],
        "extra_text": [("Note: provisional; values may be revised.", 185)],
    }])
    result = parse_pdf_result("", pdf, filename="footnote.pdf")
    assert result.tables[0].headers == ["Production (MT)", "April"]
    assert "provisional" in result.pages[0].text
    assert result.tables[0].cells[-1]["value"] == "781.05"


def test_xlsx_cells_preserve_sheet_coordinate_formula_display_and_merges(tmp_path):
    workbook = openpyxl.Workbook()
    production = workbook.active
    production.title = "Production"
    production.merge_cells("A1:B1")
    production["A1"] = "Production (MT)"
    production["A2"] = "April"
    production["B2"] = "May"
    production["A3"] = "CIL"
    production["B3"] = 781.05
    production["C3"] = "=B3*2"
    despatch = workbook.create_sheet("Despatch")
    despatch["D17"] = 78.105
    path = tmp_path / "stats.xlsx"
    workbook.save(path)

    result = parse_document_result(str(path), "XLSX", filename="stats.xlsx")
    assert {table.sheet_name for table in result.tables} == {"Production", "Despatch"}
    production_table = next(table for table in result.tables if table.sheet_name == "Production")
    cells = {cell["coordinate"]: cell for cell in production_table.cells}
    assert cells["B3"]["value"] == 781.05
    assert cells["C3"]["value"] == "=B3*2"
    assert production_table.formulas["C3"] == "=B3*2"
    assert production_table.merged_cells == ["A1:B1"]
    assert cells["C3"]["displayed_value"] in (None, "=B3*2", 1562.1)
    assert any(cell["is_blank"] for cell in production_table.cells)
    despatch_table = next(table for table in result.tables if table.sheet_name == "Despatch")
    assert next(cell for cell in despatch_table.cells if cell["coordinate"] == "D17")["value"] == 78.105


def test_csv_cells_preserve_coordinates_and_blank_values(tmp_path):
    path = tmp_path / "values.csv"
    path.write_text("Metric,April,May\nProduction,781.05,\nChange,-5,0.781\n", encoding="utf-8")
    result = parse_document_result(str(path), "CSV", filename="values.csv")
    table = result.tables[0]
    cells = {cell["coordinate"]: cell for cell in table.cells}
    assert cells["B2"]["value"] == "781.05"
    assert cells["C2"]["value"] == ""
    assert cells["A3"]["row"] == 3
    assert cells["B3"]["column"] == 2


def test_scanned_table_preserves_ocr_evidence_without_fabricating_cells():
    image = Image.new("RGB", (1400, 900), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 64)
    for xy, text in [((110, 170), "Metric"), ((530, 170), "April"), ((930, 170), "May"), ((110, 310), "Production"), ((530, 310), "781.05"), ((930, 310), "781.50")]:
        draw.text(xy, text, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    pdf_document = fitz.open()
    page = pdf_document.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=buffer.getvalue())
    result = parse_pdf_result("", pdf_document.tobytes(), filename="scanned-table.pdf")
    assert result.pages[0].classification == "SCANNED"
    assert result.pages[0].extraction_method == "OCR"
    assert "781.05" in result.pages[0].text or "781.50" in result.pages[0].text
    assert result.pages[0].blocks
    # No table rows are invented from OCR text by the generic Step 1 parser.
    assert result.tables == []


def test_corrupt_pdf_fails_explicitly_without_partial_table_result(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(Exception):
        parse_document_result(str(path), filename="broken.pdf")
