"""Generic document parsing adapters for the common Step 1 model."""

import csv
import io
import logging
import os
import re
import math
import shutil
import time
from typing import Any, Dict, List, Optional, Tuple

from app.services.document_models import DocumentResult, EvidenceBlock, ImageResult, PageResult, TableResult, page_to_legacy_dict
from app.services.document_ai_client import DocumentAIError, get_document_ai_client
from app.services.text_sanitization import sanitize_document_result
from config import settings

logger = logging.getLogger(__name__)
# A result that barely clears Tesseract's average confidence is not sufficient
# evidence for an image-only page.  Keep the threshold conservative so degraded
# fallback OCR is reviewable instead of being reported as clean extraction.
OCR_REVIEW_THRESHOLD = 0.70

try:
    import fitz
    HAS_PYMUPDF = True
except ImportError:
    fitz = None
    HAS_PYMUPDF = False

try:
    import pytesseract
    from PIL import Image, ImageFilter, ImageOps
    HAS_OCR = True
except ImportError:
    pytesseract = None
    Image = ImageFilter = ImageOps = None
    HAS_OCR = False

if HAS_OCR and pytesseract is not None and shutil.which("tesseract") is None:
    # Keep OCR usable when a Windows service/test runner has a restricted PATH.
    # Do not download or silently substitute another OCR engine.
    bundled_tesseract = os.path.join(os.environ.get("ProgramFiles", r"C:\Program Files"), "Tesseract-OCR", "tesseract.exe")
    if os.path.isfile(bundled_tesseract):
        pytesseract.pytesseract.tesseract_cmd = bundled_tesseract


def _meaningful_char_count(text: str) -> int:
    return len(re.findall(r"[\w\d]", text or "", flags=re.UNICODE))


def classify_pdf_page_signals(text_char_count: int, text_block_count: int, image_count: int, image_coverage: float) -> Tuple[str, bool, float]:
    """Classify one page independently; returns (class, needs_ocr, confidence)."""
    meaningful = text_char_count >= 20 and text_block_count > 0
    # Some digitally authored PDFs carry a full-page background, illustration,
    # or rendered design layer underneath a complete native text layer. Treat
    # that combination as DIGITAL: the native layer is the reliable source and
    # OCRing the background would be both redundant and very expensive. A
    # partial raster with native text remains MIXED and still gets selective
    # raster inspection.
    reliable_native_layer = text_char_count >= 40 and text_block_count >= 1
    if meaningful and image_count >= 1 and image_coverage >= 0.95 and reliable_native_layer:
        return "DIGITAL", False, min(0.96, 0.82 + min(text_char_count, 800) / 4000)
    large_image = image_coverage >= 0.35 or (image_count >= 1 and image_coverage >= 0.18)
    if meaningful and not large_image:
        return "DIGITAL", False, min(0.99, 0.72 + min(text_char_count, 500) / 1800)
    if meaningful and large_image:
        # Mixed pages must preserve native text and inspect raster content too.
        return "MIXED", True, 0.65
    if image_count or image_coverage > 0.05:
        return "SCANNED", True, 0.35
    return "DIGITAL", False, 0.20


def _preprocess_for_ocr(image: Any) -> Any:
    if not HAS_OCR:
        return image
    image = ImageOps.exif_transpose(image).convert("L")
    image = ImageOps.autocontrast(image)
    image = _deskew_image(image)
    image = image.filter(ImageFilter.MedianFilter(size=3))
    thresholded = image.point(lambda px: 255 if px > 185 else 0)
    return thresholded if thresholded.getbbox() else image


def _deskew_image(image: Any) -> Any:
    """Apply a small deterministic deskew correction using row-projection variance."""
    if not HAS_OCR or image.width < 40 or image.height < 40:
        return image
    scale = min(1.0, 900.0 / max(image.width, image.height))
    sample = image.resize((max(1, int(image.width * scale)), max(1, int(image.height * scale)))) if scale < 1 else image
    pixels = sample.load()

    def score(angle: float) -> float:
        rotated = sample.rotate(angle, expand=True, fillcolor=255)
        rot_pixels = rotated.load()
        counts = []
        for y in range(0, rotated.height, max(1, rotated.height // 300)):
            counts.append(sum(1 for x in range(0, rotated.width, max(1, rotated.width // 900)) if rot_pixels[x, y] < 180))
        if not counts:
            return 0.0
        mean = sum(counts) / len(counts)
        return sum((value - mean) ** 2 for value in counts) / len(counts)

    angles = (-3.0, -2.0, -1.0, 0.0, 1.0, 2.0, 3.0)
    best = max(angles, key=score)
    return image.rotate(best, expand=True, fillcolor=255) if math.fabs(best) >= 0.5 else image


def _ocr_image(image: Any, config: Optional[str] = None) -> Tuple[str, List[EvidenceBlock], Optional[float]]:
    if not HAS_OCR:
        return "", [], None
    data = pytesseract.image_to_data(image, config=config or "", output_type=pytesseract.Output.DICT)
    blocks: List[EvidenceBlock] = []
    words: List[str] = []
    confidences: List[float] = []
    for i, raw_text in enumerate(data.get("text", [])):
        text = (raw_text or "").strip()
        try:
            confidence = float(data.get("conf", ["-1"])[i])
        except (ValueError, TypeError, IndexError):
            confidence = -1.0
        if not text or confidence < 0:
            continue
        x, y = float(data.get("left", [0])[i]), float(data.get("top", [0])[i])
        w, h = float(data.get("width", [0])[i]), float(data.get("height", [0])[i])
        normalized = max(0.0, min(1.0, confidence / 100.0))
        blocks.append(EvidenceBlock(text=text, bbox=[x, y, x + w, y + h], confidence=normalized, method="OCR"))
        words.append(text)
        confidences.append(normalized)
    return " ".join(words), blocks, (sum(confidences) / len(confidences) if confidences else None)


def _has_visual_content(image: Any) -> bool:
    """Distinguish a genuinely blank raster from an engine returning EMPTY."""
    if not HAS_OCR or image is None:
        return False
    grayscale = ImageOps.grayscale(image)
    scale = min(1.0, 160.0 / max(grayscale.width, grayscale.height))
    sample = grayscale.resize((max(1, int(grayscale.width * scale)), max(1, int(grayscale.height * scale)))) if scale < 1 else grayscale
    flattened = getattr(sample, "get_flattened_data", None)
    pixels = list(flattened()) if callable(flattened) else list(sample.getdata())
    # Ignore page borders/compression speckle; readable content occupies a
    # materially larger fraction of the sampled raster.
    return bool(pixels) and sum(1 for pixel in pixels if pixel < 220) / len(pixels) > 0.01


def _image_png_bytes(image: Any) -> bytes:
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _paddle_blocks(response: Dict[str, Any]) -> List[EvidenceBlock]:
    blocks: List[EvidenceBlock] = []
    for raw in response.get("blocks") or []:
        if not isinstance(raw, dict) or not str(raw.get("text") or "").strip():
            continue
        metadata = dict(raw.get("metadata") or {})
        metadata.update({"engine": response.get("engine", "paddle_ppocrv6"), "model": response.get("model"), "version": response.get("version"), "source_raster": response.get("source_raster") or {}})
        blocks.append(EvidenceBlock(text=str(raw["text"]), bbox=raw.get("bbox"), confidence=raw.get("confidence"), method=str(raw.get("method") or response.get("engine") or "paddle_ppocrv6"), block_type=str(raw.get("block_type") or "TEXT"), metadata=metadata))
    return blocks


def _paddle_tables(response: Dict[str, Any]) -> List[TableResult]:
    tables: List[TableResult] = []
    for raw in response.get("tables") or []:
        if not isinstance(raw, dict):
            continue
        tables.append(TableResult(
            page_number=int(raw.get("page_number") or response.get("page_number") or 1),
            table_number=int(raw.get("table_number") or len(tables) + 1),
            headers=raw.get("headers") or [],
            rows=raw.get("rows") or [],
            bounding_box=raw.get("bounding_box"),
            extraction_confidence=raw.get("extraction_confidence"),
            extraction_method=str(raw.get("extraction_method") or "paddle_ppstructurev3"),
            cells=raw.get("cells") or [],
            warnings=raw.get("warnings") or [],
        ))
    return tables


def _table_candidate(text: str, blocks: List[EvidenceBlock], classification: str) -> bool:
    mode = str(settings.DOCUMENT_AI_STRUCTURE_MODE or "table_hint").lower()
    if mode == "disabled" or classification not in {"SCANNED", "MIXED"}:
        return False
    if mode == "always":
        return True
    numeric_blocks = sum(1 for block in blocks if re.search(r"\d", block.text or ""))
    row_groups: Dict[int, int] = {}
    for block in blocks:
        if block.bbox and len(block.bbox) == 4:
            row_key = round(((block.bbox[1] + block.bbox[3]) / 2) / 12)
            row_groups[row_key] = row_groups.get(row_key, 0) + 1
    # Require repeated multi-column rows. A numeric-heavy prose page often
    # has many numeric tokens but only one OCR block per line; routing those
    # pages to the expensive structure model creates false table warnings.
    multi_column_rows = sum(1 for count in row_groups.values() if count >= 2)
    return len(blocks) >= 4 and numeric_blocks >= 2 and multi_column_rows >= 2


def _tesseract_blocks(blocks: List[EvidenceBlock], *, fallback_from: Optional[str] = None) -> List[EvidenceBlock]:
    result: List[EvidenceBlock] = []
    for block in blocks:
        metadata = dict(block.metadata or {})
        metadata.update({"engine": "tesseract", "model": "tesseract-runtime", "fallback_from": fallback_from})
        result.append(EvidenceBlock(text=block.text, bbox=block.bbox, confidence=block.confidence, method="tesseract", block_type=block.block_type, metadata=metadata))
    return result


def _ocr_region_with_routing(image: Any, *, page_number: int, source_raster: Dict[str, Any], classification: str, filename: Optional[str]) -> Tuple[str, List[EvidenceBlock], Optional[float], str, Dict[str, Any], List[TableResult], List[str]]:
    """Run Paddle first, then the existing Tesseract path when recoverable."""
    warnings: List[str] = []
    paddle_failure: Optional[str] = None
    if settings.DOCUMENT_AI_ENABLED:
        try:
            client = get_document_ai_client()
            image_bytes = _image_png_bytes(image)
            response = client.ocr(image_bytes, page_number=page_number, source_raster=source_raster, filename=filename)
            paddle_blocks = _paddle_blocks(response)
            paddle_text = str(response.get("text") or "").strip()
            if paddle_text or not _has_visual_content(image):
                tables: List[TableResult] = []
                if paddle_text and settings.DOCUMENT_AI_ENABLE_STRUCTURE and _table_candidate(paddle_text, paddle_blocks, classification):
                    try:
                        structure = client.structure(image_bytes, page_number=page_number, source_raster=source_raster, filename=filename)
                        tables = _paddle_tables(structure)
                        warnings.extend(structure.get("warnings") or [])
                    except DocumentAIError as exc:
                        warnings.append(f"Page {page_number}: PP-StructureV3 unavailable; no fabricated table emitted ({str(exc)[:180]})")
                return paddle_text, paddle_blocks, response.get("confidence"), response.get("engine", "paddle_ppocrv6"), {"engine": response.get("engine", "paddle_ppocrv6"), "model": response.get("model"), "version": response.get("version"), "source_raster": source_raster, "service_status": response.get("status", "OK")}, tables, warnings
            paddle_failure = "Paddle returned EMPTY for a non-blank raster"
        except DocumentAIError as exc:
            paddle_failure = str(exc)
        if paddle_failure:
            warnings.append(f"Page {page_number}: Paddle OCR failed; Tesseract fallback used ({paddle_failure[:180]})")

    if not settings.DOCUMENT_AI_ENABLED and not settings.DOCUMENT_AI_FALLBACK_TO_TESSERACT:
        return "", [], None, "paddle_disabled", {"engine": "none", "service_status": "disabled"}, [], warnings
    try:
        text, blocks, confidence = _ocr_image(image)
    except Exception as exc:
        warnings.append(f"Page {page_number}: Tesseract fallback failed: {type(exc).__name__}: {str(exc)[:180]}")
        return "", [], 0.0, "tesseract", {"engine": "tesseract", "fallback_from": paddle_failure, "service_status": "failed"}, [], warnings
    return text, _tesseract_blocks(blocks, fallback_from="paddle_ppocrv6" if paddle_failure else None), confidence, "tesseract", {"engine": "tesseract", "model": "tesseract-runtime", "fallback_from": "paddle_ppocrv6" if paddle_failure else None, "service_status": "fallback" if paddle_failure else "disabled"}, [], warnings


def _embedded_page_images(document: Any, page: Any) -> List[Any]:
    """Return embedded page rasters for OCR fallback without inventing geometry."""
    if not HAS_OCR:
        return []
    images = []
    for image_info in page.get_images(full=True):
        try:
            image_bytes = document.extract_image(image_info[0])["image"]
            with Image.open(io.BytesIO(image_bytes)) as embedded:
                images.append(ImageOps.exif_transpose(embedded).convert("RGB"))
        except Exception as exc:
            logger.debug("Embedded page image unavailable for OCR fallback: %s", exc)
    return images


def _ocr_pdf_page(document: Any, page: Any, rendered_image: Any, *, page_number: int = 1, classification: str = "SCANNED", filename: Optional[str] = None) -> Tuple[str, List[EvidenceBlock], Optional[float], str, Dict[str, Any], List[TableResult], List[str]]:
    """Route one page/region through Paddle, with explicit Tesseract recovery."""
    processed = _preprocess_for_ocr(rendered_image)
    regions = []
    if classification == "MIXED":
        for index, embedded in enumerate(_embedded_page_images(document, page), start=1):
            regions.append((embedded, {"kind": "embedded_image", "image_number": index, "width": embedded.width, "height": embedded.height}))
    if not regions:
        regions = [(processed, {"kind": "rendered_page", "width": processed.width, "height": processed.height, "dpi": 300, "preprocessing": ["grayscale", "deskew", "denoise", "autocontrast", "thresholding"]})]

    text_parts: List[str] = []
    blocks: List[EvidenceBlock] = []
    confidences: List[float] = []
    sources: List[str] = []
    metadata: Dict[str, Any] = {"regions": [], "fallbacks": []}
    tables: List[TableResult] = []
    warnings: List[str] = []
    for image, source_raster in regions:
        region_text, region_blocks, region_confidence, engine, region_metadata, region_tables, region_warnings = _ocr_region_with_routing(image, page_number=page_number, source_raster=source_raster, classification=classification, filename=filename)
        if region_text.strip():
            text_parts.append(region_text.strip())
        blocks.extend(region_blocks)
        if region_confidence is not None:
            confidences.append(float(region_confidence))
        sources.append(engine)
        metadata["regions"].append(region_metadata)
        if region_metadata.get("fallback_from"):
            metadata["fallbacks"].append(region_metadata)
        tables.extend(region_tables)
        warnings.extend(region_warnings)
    # Preserve the Step-1 screenshot recovery path when the Paddle request
    # failed or returned no usable text for a rendered page. The embedded
    # raster has different pixels from the preprocessed 300-DPI page render.
    if not text_parts and regions and regions[0][1].get("kind") == "rendered_page":
        for index, embedded in enumerate(_embedded_page_images(document, page), start=1):
            try:
                try:
                    fallback_text, fallback_blocks, fallback_confidence = _ocr_image(embedded, config="--psm 11")
                except TypeError:
                    # Compatibility with deterministic test seams that accept
                    # only the historical one-argument helper.
                    fallback_text, fallback_blocks, fallback_confidence = _ocr_image(embedded)
            except Exception as exc:
                warnings.append(f"Page {page_number}: embedded-raster Tesseract recovery failed: {type(exc).__name__}: {str(exc)[:180]}")
                continue
            if fallback_text.strip():
                text_parts.append(fallback_text.strip())
                blocks = _tesseract_blocks(fallback_blocks, fallback_from="paddle_ppocrv6")
                confidences = [float(fallback_confidence)] if fallback_confidence is not None else []
                source = "embedded_raster"
                sources.append("embedded_raster")
                metadata["engine"] = "tesseract"
                metadata["model"] = "tesseract-runtime"
                metadata["fallback_from"] = "paddle_ppocrv6"
                metadata["regions"].append({"engine": "tesseract", "model": "tesseract-runtime", "fallback_from": "paddle_ppocrv6", "source_raster": {"kind": "embedded_image", "image_number": index, "width": embedded.width, "height": embedded.height}})
                metadata["fallbacks"].append(metadata["regions"][-1])
                warnings.append(f"Page {page_number}: embedded-raster Tesseract recovery used; review recommended")
                break
    unique_sources = list(dict.fromkeys(sources))
    source = "embedded_raster" if "embedded_raster" in unique_sources else unique_sources[0] if len(unique_sources) == 1 else "+".join(unique_sources)
    metadata["engine"] = source
    metadata["model"] = next((item.get("model") for item in metadata["regions"] if item.get("model")), None)
    metadata["fallback_from"] = next((item.get("fallback_from") for item in metadata["regions"] if item.get("fallback_from")), None)
    if source == "tesseract":
        source = "embedded_raster" if any(item.get("kind") == "embedded_image" for _, item in regions) else "rendered_300dpi_preprocessed"
    return "\n".join(text_parts), blocks, (sum(confidences) / len(confidences) if confidences else None), source, metadata, tables, warnings


def _native_blocks(page: Any) -> Tuple[str, List[EvidenceBlock], int, float]:
    page_dict = page.get_text("dict")
    blocks: List[EvidenceBlock] = []
    texts: List[str] = []
    text_blocks = 0
    image_area = 0.0
    for block in page_dict.get("blocks", []):
        bbox = list(block.get("bbox", [])) or None
        if block.get("type") == 0:
            text = "\n".join(span.get("text", "") for line in block.get("lines", []) for span in line.get("spans", [])).strip()
            if text:
                texts.append(text)
                text_blocks += 1
                blocks.append(EvidenceBlock(text=text, bbox=bbox, confidence=0.98, method="NATIVE"))
        elif block.get("type") == 1 and bbox and len(bbox) == 4:
            image_area += max(0.0, bbox[2] - bbox[0]) * max(0.0, bbox[3] - bbox[1])
    page_area = max(1.0, float(page.rect.width) * float(page.rect.height))
    return "\n".join(texts).strip(), blocks, text_blocks, min(1.0, image_area / page_area)


def _chart_like_drawing_count(page: Any, table_bbox: Optional[List[float]]) -> int:
    """Count filled vector regions inside a candidate table bounding box.

    Filled bars/areas are strong figure evidence when a table detector has
    also produced a very sparse, oversized grid.  A normal sparse table with
    no chart-like drawing evidence is intentionally left eligible for table
    persistence.
    """
    if not table_bbox or len(table_bbox) != 4 or not hasattr(page, "get_drawings"):
        return 0
    try:
        candidate = tuple(float(value) for value in table_bbox)
        x1, y1, x2, y2 = candidate
        count = 0
        for drawing in page.get_drawings() or []:
            fill = drawing.get("fill")
            rect = drawing.get("rect")
            if not fill or rect is None:
                continue
            dx1, dy1, dx2, dy2 = float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)
            intersection_width = max(0.0, min(x2, dx2) - max(x1, dx1))
            intersection_height = max(0.0, min(y2, dy2) - max(y1, dy1))
            if intersection_width * intersection_height > 4.0:
                count += 1
        return count
    except Exception:
        # Table classification must remain fail-safe if a PDF drawing object
        # is malformed or unsupported by the installed PyMuPDF version.
        return 0


def _reject_figure_like_table_candidate(page: Any, raw_rows: List[List[Any]], table: Any) -> Optional[Dict[str, Any]]:
    """Return a rejection record for a clearly chart-like native table grid.

    PyMuPDF can mistake chart axes/bars and their numeric annotations for a
    huge table.  Do not reject sparse tables on sparsity alone: require both a
    large, very low-occupancy grid and multiple filled vector regions in its
    bounding box.  The latter is the evidence that distinguishes the known
    bar-chart failure from a legitimate sparse text table.
    """
    row_count = len(raw_rows)
    column_count = max((len(row) for row in raw_rows), default=0)
    if row_count < 5 or column_count < 8:
        return None
    non_empty = sum(
        1
        for row in raw_rows
        for value in row
        if value is not None and str(value).strip()
    )
    occupancy = non_empty / max(1, row_count * column_count)
    if occupancy > 0.20:
        return None
    bbox = list(getattr(table, "bbox", []) or []) or None
    filled_drawing_count = _chart_like_drawing_count(page, bbox)
    if filled_drawing_count < 3:
        return None
    return {
        "classification": "FIGURE_OR_CHART",
        "reason": "sparse oversized native grid with filled vector regions",
        "bounding_box": bbox,
        "row_count": row_count,
        "column_count": column_count,
        "non_empty_cell_count": non_empty,
        "occupancy": round(occupancy, 4),
        "filled_vector_region_count": filled_drawing_count,
        "extraction_method": "NATIVE_TABLE",
    }


def _extract_pdf_tables(page: Any, page_number: int) -> Tuple[List[TableResult], List[str], List[Dict[str, Any]]]:
    tables: List[TableResult] = []
    warnings: List[str] = []
    rejected_candidates: List[Dict[str, Any]] = []
    if not hasattr(page, "find_tables"):
        return tables, warnings, rejected_candidates
    found = None
    used_text_strategy = False
    try:
        found = page.find_tables()
    except Exception as exc:
        warnings.append(f"Native table detection warning: {type(exc).__name__}: {str(exc)[:180]}")

    # PyMuPDF's default detector is strongest for ruled tables.  A second,
    # bounded text-strategy pass recovers simple borderless tables without
    # making the parser depend on a report-specific layout.
    if not getattr(found, "tables", None):
        try:
            text_found = page.find_tables(strategy="text")
            if getattr(text_found, "tables", None):
                found = text_found
                used_text_strategy = True
        except (TypeError, ValueError) as exc:
            warnings.append(f"Borderless table detection unavailable: {type(exc).__name__}: {str(exc)[:180]}")
        except Exception as exc:
            warnings.append(f"Borderless table detection warning: {type(exc).__name__}: {str(exc)[:180]}")

    try:
        for index, table in enumerate(getattr(found, "tables", []) or [], start=1):
            raw_rows = table.extract() or []
            if not raw_rows:
                continue
            rows = [["" if value is None else value for value in row] for row in raw_rows]
            rejection = _reject_figure_like_table_candidate(page, raw_rows, table)
            if rejection is not None:
                rejection["candidate_number"] = index
                rejection["page_number"] = page_number
                rejected_candidates.append(rejection)
                warnings.append(
                    f"Native table candidate {index} rejected as likely figure/chart "
                    f"(page {page_number}, bbox={rejection['bounding_box']}, "
                    f"occupancy={rejection['occupancy']:.3f}); source text/evidence preserved"
                )
                continue
            headers = rows[0] if rows else []
            table_rows = rows[1:] if len(rows) > 1 else []
            table_warnings = [] if table_rows else ["Table contains no data rows"]
            if used_text_strategy:
                table_warnings.append("Borderless table detected with text strategy")

            # PyMuPDF exposes table.cells in column-major order.  Keep the
            # source row/column indexes and only attach a bbox when the
            # library supplied one; never invent geometry for merged/unknown
            # cells.
            cell_records: List[Dict[str, Any]] = []
            raw_cells = list(getattr(table, "cells", []) or [])
            row_count = len(rows)
            col_count = max((len(row) for row in rows), default=0)
            missing_geometry = False
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    bbox = None
                    cell_index = column_index * row_count + row_index
                    if cell_index < len(raw_cells) and raw_cells[cell_index]:
                        try:
                            bbox = list(raw_cells[cell_index])
                        except TypeError:
                            bbox = None
                    if bbox is None:
                        missing_geometry = True
                    cell_records.append({
                        "row_index": row_index,
                        "column_index": column_index,
                        "source_row": row_index + 1,
                        "source_column": column_index + 1,
                        "value": value,
                        "raw_value": value,
                        "bounding_box": bbox,
                        "bbox": bbox,
                        "is_header": row_index == 0,
                        "extraction_method": "NATIVE_TABLE",
                    })
            if missing_geometry:
                table_warnings.append("Cell geometry unavailable for one or more cells")
            if not cell_records:
                table_warnings.append("No cell-level provenance was returned by the table detector")
            confidence = 0.90 if len(headers) > 1 and table_rows else 0.55
            if used_text_strategy:
                confidence = min(confidence, 0.80)
            if missing_geometry:
                confidence = min(confidence, 0.75)
            tables.append(TableResult(
                page_number=page_number,
                # Number persisted tables, not rejected detector candidates;
                # candidate_number remains in rejection metadata for audit.
                table_number=len(tables) + 1,
                headers=headers,
                rows=table_rows,
                bounding_box=list(getattr(table, "bbox", []) or []) or None,
                extraction_confidence=confidence,
                extraction_method="NATIVE_TABLE",
                cells=cell_records,
                warnings=table_warnings,
            ))
    except Exception as exc:
        warnings.append(f"Native table extraction warning: {type(exc).__name__}: {str(exc)[:180]}")
    return tables, warnings, rejected_candidates


def parse_pdf_result(file_path: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None) -> DocumentResult:
    if not HAS_PYMUPDF:
        raise RuntimeError("PyMuPDF is required for PDF ingestion")
    doc = fitz.open(stream=file_bytes, filetype="pdf") if file_bytes is not None else fitz.open(file_path)
    result = DocumentResult(document_id=None, filename=filename or os.path.basename(file_path), file_type="PDF", extraction_method="NATIVE")
    result.metadata.update({"page_count": len(doc), "parser": "PyMuPDF", "timings_ms": {"pages": []}})
    try:
        for page_index in range(len(doc)):
            page_started = time.perf_counter()
            page_timing: Dict[str, float] = {}
            page_number = page_index + 1
            stage_started = time.perf_counter()
            page = doc.load_page(page_index)
            page_timing["page_load_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            stage_started = time.perf_counter()
            native_text, native_blocks, block_count, image_coverage = _native_blocks(page)
            page_timing["native_text_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            stage_started = time.perf_counter()
            image_count = len(page.get_images(full=True))
            source_image_dimensions = []
            for image_info in page.get_images(full=True):
                try:
                    embedded = fitz.Pixmap(doc, image_info[0])
                    source_image_dimensions.append([embedded.width, embedded.height])
                except Exception:
                    continue
            page_timing["raster_inspection_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            stage_started = time.perf_counter()
            classification, requires_ocr, base_confidence = classify_pdf_page_signals(_meaningful_char_count(native_text), block_count, image_count, image_coverage)
            page_timing["classification_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            final_text, blocks, method, confidence = native_text, native_blocks, "NATIVE", base_confidence
            ocr_source = None
            ocr_metadata: Dict[str, Any] = {}
            ai_tables: List[TableResult] = []
            page_timing["ocr_ms"] = 0.0
            if requires_ocr:
                if HAS_OCR:
                    stage_started = time.perf_counter()
                    pix = page.get_pixmap(dpi=300, alpha=False)
                    image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                    try:
                        ocr_text, ocr_blocks, ocr_confidence, ocr_source, ocr_metadata, ai_tables, routing_warnings = _ocr_pdf_page(doc, page, image, page_number=page_number, classification=classification, filename=filename or os.path.basename(file_path))
                        result.warnings.extend(routing_warnings)
                    except Exception as exc:
                        ocr_text, ocr_blocks, ocr_confidence, routing_warnings = "", [], 0.0, []
                        ocr_source = "error"
                        ocr_metadata = {"engine": "unknown", "service_status": "failed"}
                        result.warnings.append(f"Page {page_number}: OCR failed: {type(exc).__name__}: {str(exc)[:180]}")
                    if ocr_text:
                        final_text = (native_text + "\n" + ocr_text).strip() if classification == "MIXED" and native_text else ocr_text
                        blocks = native_blocks + ocr_blocks if classification == "MIXED" else ocr_blocks
                        method = "NATIVE+OCR" if classification == "MIXED" else "OCR"
                        confidence = ocr_confidence
                    else:
                        method = "NATIVE+OCR" if classification == "MIXED" else "OCR"
                        confidence = ocr_confidence if ocr_confidence is not None else 0.0
                        if ocr_metadata.get("service_status") != "EMPTY" or _has_visual_content(image):
                            result.warnings.append(f"Page {page_number}: OCR returned no text")
                    if ocr_source and ocr_source not in {"rendered_300dpi_preprocessed", "paddle_ppocrv6"}:
                        result.warnings.append(f"Page {page_number}: OCR fallback used from {ocr_source}; review recommended")
                    if ocr_confidence is not None and ocr_confidence < OCR_REVIEW_THRESHOLD:
                        result.warnings.append(f"Page {page_number}: low OCR confidence ({ocr_confidence:.2f}); review recommended")
                    page_timing["ocr_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
                else:
                    result.warnings.append(f"Page {page_number}: OCR required but Tesseract/Pillow is unavailable")
            if not final_text:
                result.warnings.append(f"Page {page_number}: no readable text extracted")
            if requires_ocr and source_image_dimensions and min(min(dimensions) for dimensions in source_image_dimensions) < 500:
                result.warnings.append(f"Page {page_number}: low-resolution source image; review recommended")
            stage_started = time.perf_counter()
            page_tables, table_warnings, rejected_table_candidates = _extract_pdf_tables(page, page_number)
            page_timing["native_table_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            if ai_tables and not page_tables:
                page_tables = ai_tables
            result.tables.extend(page_tables)
            result.warnings.extend(f"Page {page_number}: {warning}" for warning in table_warnings)
            stage_started = time.perf_counter()
            page_timing["normalization_ms"] = round((time.perf_counter() - stage_started) * 1000, 2)
            page_timing["page_total_ms"] = round((time.perf_counter() - page_started) * 1000, 2)
            page_metadata = {"native_text_char_count": len(native_text), "meaningful_text_char_count": _meaningful_char_count(native_text), "native_text_block_count": block_count, "image_count": image_count, "image_coverage": image_coverage, "ocr_requested": requires_ocr, "ocr_source": ocr_source, "ocr_engine": ocr_metadata.get("engine"), "ocr_model": ocr_metadata.get("model"), "ocr_version": ocr_metadata.get("version"), "ocr_service_status": ocr_metadata.get("service_status"), "ocr_fallback_from": ocr_metadata.get("fallback_from"), "ocr_regions": ocr_metadata.get("regions", []), "rejected_table_candidates": rejected_table_candidates, "ocr_coordinate_space": "embedded_image_pixels" if ocr_source == "embedded_raster" else "rendered_page_pixels" if requires_ocr else None, "preprocessing": ["300dpi", "grayscale", "deskew", "denoise", "autocontrast", "thresholding"] if requires_ocr else [], "timings_ms": page_timing}
            result.pages.append(PageResult(page_number=page_number, text=final_text, extraction_method=method, confidence=confidence, classification=classification, blocks=blocks, width=float(page.rect.width), height=float(page.rect.height), metadata=page_metadata))
            result.metadata["timings_ms"]["pages"].append({"page_number": page_number, "classification": classification, "ocr_requested": requires_ocr, **page_timing})
            logger.info("PDF page timing filename=%s page=%s classification=%s ocr_requested=%s timings_ms=%s", filename or os.path.basename(file_path), page_number, classification, requires_ocr, page_timing)
    finally:
        doc.close()
    result.extraction_confidence = sum((p.confidence or 0.0) for p in result.pages) / len(result.pages) if result.pages else 0.0
    result.extraction_method = "MIXED" if any(p.extraction_method != "NATIVE" for p in result.pages) else "NATIVE"
    return result


def parse_image_result(file_path: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None, file_type: str = "IMAGE") -> DocumentResult:
    if not HAS_OCR:
        raise RuntimeError("Pillow and pytesseract are required for image ingestion")
    image = Image.open(io.BytesIO(file_bytes) if file_bytes is not None else file_path)
    original = image.copy()
    ocr_error = None
    routing_warnings: List[str] = []
    ocr_source = "tesseract"
    ocr_metadata: Dict[str, Any] = {}
    tables: List[TableResult] = []
    try:
        text, blocks, confidence, ocr_source, ocr_metadata, tables, routing_warnings = _ocr_region_with_routing(_preprocess_for_ocr(original), page_number=1, source_raster={"kind": "standalone_image", "width": original.width, "height": original.height}, classification="SCANNED", filename=filename or os.path.basename(file_path))
    except Exception as exc:
        ocr_error = exc
        text, blocks, confidence = "", [], 0.0
    result = DocumentResult(document_id=None, filename=filename or os.path.basename(file_path), file_type=file_type.upper(), extraction_method="OCR", extraction_confidence=confidence)
    result.tables.extend(tables)
    result.pages.append(PageResult(page_number=1, text=text, extraction_method="OCR", confidence=confidence, classification="SCANNED", blocks=blocks, width=float(original.width), height=float(original.height), metadata={"preprocessing": ["grayscale", "deskew", "denoise", "autocontrast", "thresholding"], "ocr_source": ocr_source, "ocr_engine": ocr_metadata.get("engine"), "ocr_model": ocr_metadata.get("model"), "ocr_version": ocr_metadata.get("version"), "ocr_service_status": ocr_metadata.get("service_status"), "ocr_fallback_from": ocr_metadata.get("fallback_from"), "ocr_regions": ocr_metadata.get("regions", [])}))
    result.images.append(ImageResult(image_number=1, page_number=1, source="standalone", mime_type=Image.MIME.get(getattr(original, "format", ""), None), width=original.width, height=original.height, text=text, ocr_confidence=confidence, metadata={"format": original.format, "mode": original.mode, "ocr_engine": ocr_metadata.get("engine"), "ocr_model": ocr_metadata.get("model"), "ocr_source": ocr_source}))
    result.warnings.extend(routing_warnings)
    if not text:
        result.warnings.append("Image OCR returned no text")
    if ocr_error is not None:
        result.warnings.append(f"Image OCR failed: {type(ocr_error).__name__}: {str(ocr_error)[:180]}")
    if min(original.width, original.height) < 500:
        result.warnings.append("Image source resolution is low; review recommended")
    if confidence is not None and confidence < OCR_REVIEW_THRESHOLD:
        result.warnings.append(f"Image: low OCR confidence ({confidence:.2f}); review recommended")
    return result


def parse_docx_result(file_path: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None) -> DocumentResult:
    import docx
    document = docx.Document(io.BytesIO(file_bytes) if file_bytes is not None else file_path)
    result = DocumentResult(document_id=None, filename=filename or os.path.basename(file_path), file_type="DOCX", extraction_method="NATIVE", extraction_confidence=0.98)
    result.title = document.core_properties.title or None
    result.metadata.update({"author": document.core_properties.author or None, "subject": document.core_properties.subject or None, "paragraph_count": len(document.paragraphs), "table_count": len(document.tables)})
    image_number = 0
    for relationship in document.part.rels.values():
        if "image" not in relationship.reltype:
            continue
        image_number += 1
        blob = getattr(relationship.target_part, "blob", b"")
        width = height = None
        if HAS_OCR and blob:
            try:
                with Image.open(io.BytesIO(blob)) as embedded:
                    width, height = embedded.width, embedded.height
            except Exception:
                pass
        result.images.append(ImageResult(image_number=image_number, source="embedded", mime_type=getattr(relationship.target_part, "content_type", None), width=width, height=height, metadata={"relationship_id": relationship.rId, "part_name": str(getattr(relationship.target_part, "partname", ""))}))
    ordered, blocks, table_index = [], [], 0
    for element in document.element.body.iterchildren():
        if element.tag.endswith("}p"):
            paragraph = next((p for p in document.paragraphs if p._p is element), None)
            if paragraph and paragraph.text.strip():
                ordered.append(paragraph.text.strip())
                blocks.append(EvidenceBlock(text=paragraph.text.strip(), method="NATIVE", block_type="PARAGRAPH", metadata={"style": paragraph.style.name if paragraph.style else None}))
        elif element.tag.endswith("}tbl"):
            table = next((t for t in document.tables if t._tbl is element), None)
            if table is not None:
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                table_index += 1
                result.tables.append(TableResult(page_number=1, table_number=table_index, headers=rows[0] if rows else [], rows=rows[1:] if len(rows) > 1 else [], extraction_confidence=0.95, extraction_method="NATIVE_DOCX"))
                table_text = "\n".join(" | ".join(row) for row in rows)
                ordered.append(table_text)
                blocks.append(EvidenceBlock(text=table_text, method="NATIVE", block_type="TABLE", metadata={"table_number": table_index}))
    result.pages.append(PageResult(page_number=1, text="\n\n".join(ordered), extraction_method="NATIVE", confidence=0.98, blocks=blocks, metadata={"headers": [p.text for section in document.sections for p in section.header.paragraphs if p.text.strip()], "footers": [p.text for section in document.sections for p in section.footer.paragraphs if p.text.strip()]}))
    return result


def _cell_value(value: Any) -> Any:
    return value.isoformat() if hasattr(value, "isoformat") else value


def parse_xlsx_result(file_path: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None) -> DocumentResult:
    import openpyxl
    raw = file_bytes if file_bytes is not None else open(file_path, "rb").read()
    formulas_wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
    values_wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    result = DocumentResult(document_id=None, filename=filename or os.path.basename(file_path), file_type="XLSX", extraction_method="NATIVE", extraction_confidence=0.99)
    result.metadata["sheet_names"] = list(formulas_wb.sheetnames)
    for page_number, sheet_name in enumerate(formulas_wb.sheetnames, start=1):
        ws, values_ws = formulas_wb[sheet_name], values_wb[sheet_name]
        rows, cells, formulas, displayed = [], [], {}, {}
        for row in ws.iter_rows():
            row_values = []
            for cell in row:
                formula_value, displayed_value = _cell_value(cell.value), _cell_value(values_ws[cell.coordinate].value)
                row_values.append(displayed_value if displayed_value is not None else formula_value)
                cells.append({"coordinate": cell.coordinate, "row": cell.row, "column": cell.column, "value": formula_value, "displayed_value": displayed_value, "data_type": cell.data_type, "is_blank": formula_value is None})
                if cell.data_type == "f":
                    formulas[cell.coordinate], displayed[cell.coordinate] = formula_value, displayed_value
            rows.append(row_values)
        while rows and not any(value not in (None, "") for value in rows[-1]):
            rows.pop()
        merged = [str(rng) for rng in ws.merged_cells.ranges]
        table = TableResult(page_number=page_number, table_number=page_number, headers=rows[0] if rows else [], rows=rows[1:] if len(rows) > 1 else [], extraction_confidence=0.99, extraction_method="NATIVE_XLSX", sheet_name=sheet_name, cells=cells, merged_cells=merged, formulas=formulas, displayed_values=displayed)
        result.tables.append(table)
        result.pages.append(PageResult(page_number=page_number, text="\n".join(" | ".join("" if value is None else str(value) for value in row) for row in rows), extraction_method="NATIVE", confidence=0.99, metadata={"sheet_name": sheet_name, "merged_cells": merged, "cell_count": len(cells)}))
    return result


def parse_csv_result(file_path: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None) -> DocumentResult:
    raw = file_bytes if file_bytes is not None else open(file_path, "rb").read()
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig", errors="replace"))))
    def column_name(number: int) -> str:
        name = ""
        while number:
            number, remainder = divmod(number - 1, 26)
            name = chr(65 + remainder) + name
        return name
    cells = [{"coordinate": f"{column_name(column + 1)}{row_index + 1}", "row": row_index + 1, "column": column + 1, "value": value} for row_index, row in enumerate(rows) for column, value in enumerate(row)]
    result = DocumentResult(document_id=None, filename=filename or os.path.basename(file_path), file_type="CSV", extraction_method="NATIVE", extraction_confidence=0.99)
    result.tables.append(TableResult(page_number=1, table_number=1, headers=rows[0] if rows else [], rows=rows[1:] if len(rows) > 1 else [], cells=cells, extraction_confidence=0.99, extraction_method="NATIVE_CSV"))
    result.pages.append(PageResult(page_number=1, text="\n".join(" | ".join(row) for row in rows), extraction_method="NATIVE", confidence=0.99, metadata={"row_count": len(rows), "column_count": max((len(row) for row in rows), default=0)}))
    return result


def parse_document_result(file_path: str, file_type: str, file_bytes: Optional[bytes] = None, *, filename: Optional[str] = None) -> DocumentResult:
    normalized = file_type.upper().lstrip(".")
    if normalized == "PDF":
        result = parse_pdf_result(file_path, file_bytes, filename=filename)
    elif normalized == "DOCX":
        result = parse_docx_result(file_path, file_bytes, filename=filename)
    elif normalized == "XLSX":
        result = parse_xlsx_result(file_path, file_bytes, filename=filename)
    elif normalized == "CSV":
        result = parse_csv_result(file_path, file_bytes, filename=filename)
    elif normalized in {"PNG", "JPG", "JPEG", "TIF", "TIFF", "IMAGE"}:
        result = parse_image_result(file_path, file_bytes, filename=filename, file_type=normalized)
    else:
        raise ValueError(f"Unsupported document type: {file_type}")

    # All parser/OCR adapters converge here before persistence, chunking,
    # metric extraction, or vector indexing. This protects the full pipeline
    # from database-unsafe control characters without changing valid Unicode,
    # whitespace, numeric values, or geometry provenance.
    return sanitize_document_result(result)


# Compatibility shims used by the existing pipeline and tests.
def parse_pdf_document(file_path: str, file_bytes: Optional[bytes] = None) -> List[Dict[str, Any]]:
    result = parse_pdf_result(file_path, file_bytes)
    return [page_to_legacy_dict(page, result.tables) for page in result.pages]


def parse_docx_document(file_path: str, file_bytes: Optional[bytes] = None) -> List[Dict[str, Any]]:
    result = parse_docx_result(file_path, file_bytes)
    return [page_to_legacy_dict(page, result.tables) for page in result.pages]


def parse_excel_csv_document(file_path: str, file_type: str, file_bytes: Optional[bytes] = None) -> List[Dict[str, Any]]:
    result = parse_xlsx_result(file_path, file_bytes) if file_type.upper() == "XLSX" else parse_csv_result(file_path, file_bytes)
    return [page_to_legacy_dict(page, result.tables) for page in result.pages]


def parse_document_file(file_path: str, file_type: str, file_bytes: Optional[bytes] = None) -> List[Dict[str, Any]]:
    result = parse_document_result(file_path, file_type, file_bytes)
    return [page_to_legacy_dict(page, result.tables) for page in result.pages]
